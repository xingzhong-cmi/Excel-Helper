"""
File management module for handling Excel files with versioning and user isolation.
"""
import os
import shutil
from pathlib import Path
from typing import List, Dict, Optional
from datetime import datetime
import openpyxl
import pandas as pd


class FileManager:
    """Manages Excel files with versioning and user isolation."""
    
    def __init__(self, user_id: int, max_revisions: int = 10):
        """Initialize the file manager.
        
        Args:
            user_id: User ID for file isolation
            max_revisions: Maximum number of revisions to keep per file
        """
        self.user_id = user_id
        self.max_revisions = max_revisions
        self.base_path = Path(f"data/{user_id}")
        self.base_path.mkdir(parents=True, exist_ok=True)
    
    def get_user_path(self, file_id: str) -> Path:
        """Get the directory path for a file.
        
        Args:
            file_id: File identifier
            
        Returns:
            Path to the file directory
        """
        path = self.base_path / file_id
        path.mkdir(parents=True, exist_ok=True)
        return path
    
    def save_uploaded_file(self, uploaded_file, file_id: str) -> str:
        """Save an uploaded file.
        
        Args:
            uploaded_file: Streamlit uploaded file object
            file_id: File identifier
            
        Returns:
            Path to the saved file
        """
        file_path = self.get_user_path(file_id)
        revision = self._get_next_revision(file_id)
        save_path = file_path / f"rev_{revision}.xlsx"
        
        with open(save_path, "wb") as f:
            f.write(uploaded_file.getbuffer())
        
        # Save metadata
        self._save_metadata(file_id, revision, uploaded_file.name)
        
        # Clean old revisions
        self._cleanup_old_revisions(file_id)
        
        return str(save_path)
    
    def _get_next_revision(self, file_id: str) -> int:
        """Get the next revision number for a file.
        
        Args:
            file_id: File identifier
            
        Returns:
            Next revision number
        """
        file_path = self.get_user_path(file_id)
        revisions = [
            int(f.stem.split("_")[1])
            for f in file_path.glob("rev_*.xlsx")
        ]
        return max(revisions, default=0) + 1
    
    def _save_metadata(self, file_id: str, revision: int, original_name: str):
        """Save metadata for a file revision.
        
        Args:
            file_id: File identifier
            revision: Revision number
            original_name: Original filename
        """
        file_path = self.get_user_path(file_id)
        metadata_path = file_path / f"rev_{revision}_meta.txt"
        
        with open(metadata_path, "w") as f:
            f.write(f"original_name: {original_name}\n")
            f.write(f"timestamp: {datetime.now().isoformat()}\n")
    
    def _cleanup_old_revisions(self, file_id: str):
        """Remove old revisions beyond the max limit.
        
        Args:
            file_id: File identifier
        """
        file_path = self.get_user_path(file_id)
        revisions = sorted([
            int(f.stem.split("_")[1])
            for f in file_path.glob("rev_*.xlsx")
        ])
        
        if len(revisions) > self.max_revisions:
            for rev in revisions[:-self.max_revisions]:
                (file_path / f"rev_{rev}.xlsx").unlink(missing_ok=True)
                (file_path / f"rev_{rev}_meta.txt").unlink(missing_ok=True)
    
    def get_latest_revision(self, file_id: str) -> Optional[int]:
        """Get the latest revision number for a file.
        
        Args:
            file_id: File identifier
            
        Returns:
            Latest revision number or None
        """
        file_path = self.get_user_path(file_id)
        revisions = [
            int(f.stem.split("_")[1])
            for f in file_path.glob("rev_*.xlsx")
        ]
        return max(revisions) if revisions else None
    
    def get_file_path(self, file_id: str, revision: Optional[int] = None) -> Optional[str]:
        """Get the path to a specific file revision.
        
        Args:
            file_id: File identifier
            revision: Revision number (None for latest)
            
        Returns:
            Path to the file or None if not found
        """
        if revision is None:
            revision = self.get_latest_revision(file_id)
        
        if revision is None:
            return None
        
        file_path = self.get_user_path(file_id) / f"rev_{revision}.xlsx"
        return str(file_path) if file_path.exists() else None
    
    def list_revisions(self, file_id: str) -> List[Dict]:
        """List all revisions for a file.
        
        Args:
            file_id: File identifier
            
        Returns:
            List of revision info dicts
        """
        file_path = self.get_user_path(file_id)
        revisions = []
        
        for xlsx_file in sorted(file_path.glob("rev_*.xlsx")):
            rev = int(xlsx_file.stem.split("_")[1])
            meta_path = file_path / f"rev_{rev}_meta.txt"
            
            metadata = {}
            if meta_path.exists():
                with open(meta_path, "r") as f:
                    for line in f:
                        if ":" in line:
                            key, value = line.strip().split(":", 1)
                            metadata[key.strip()] = value.strip()
            
            revisions.append({
                "revision": rev,
                "path": str(xlsx_file),
                "original_name": metadata.get("original_name", f"rev_{rev}.xlsx"),
                "timestamp": metadata.get("timestamp", "")
            })
        
        return sorted(revisions, key=lambda x: x["revision"], reverse=True)
    
    def get_sheet_names(self, file_path: str) -> List[str]:
        """Get sheet names from an Excel file.
        
        Args:
            file_path: Path to the Excel file
            
        Returns:
            List of sheet names
        """
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            sheet_names = wb.sheetnames
            wb.close()
            return sheet_names
        except:
            return []
    
    def read_sheet_data(self, file_path: str, sheet_name: str, max_rows: int = 300) -> pd.DataFrame:
        """Read data from a specific sheet.
        
        Args:
            file_path: Path to the Excel file
            sheet_name: Name of the sheet
            max_rows: Maximum number of rows to read
            
        Returns:
            DataFrame containing the sheet data
        """
        try:
            df = pd.read_excel(file_path, sheet_name=sheet_name, nrows=max_rows)
            return df
        except:
            return pd.DataFrame()
    
    def get_file_metadata(self, file_path: str) -> Dict:
        """Extract metadata from an Excel file.
        
        Args:
            file_path: Path to the Excel file
            
        Returns:
            Dictionary containing file metadata
        """
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            metadata = {
                "sheets": wb.sheetnames,
                "sheet_count": len(wb.sheetnames)
            }
            
            # Get column names for each sheet
            sheet_columns = {}
            for sheet_name in wb.sheetnames[:5]:  # Limit to first 5 sheets
                ws = wb[sheet_name]
                headers = []
                for cell in ws[1]:
                    if cell.value:
                        headers.append(str(cell.value))
                sheet_columns[sheet_name] = headers[:50]  # Limit to first 50 columns
            
            metadata["columns"] = sheet_columns
            wb.close()
            
            return metadata
        except:
            return {"sheets": [], "sheet_count": 0, "columns": {}}
    
    def save_modified_file(self, file_id: str, workbook: openpyxl.Workbook) -> int:
        """Save a modified workbook as a new revision.
        
        Args:
            file_id: File identifier
            workbook: Modified workbook object
            
        Returns:
            New revision number
        """
        file_path = self.get_user_path(file_id)
        revision = self._get_next_revision(file_id)
        save_path = file_path / f"rev_{revision}.xlsx"
        
        workbook.save(save_path)
        
        # Save metadata
        with open(file_path / f"rev_{revision}_meta.txt", "w") as f:
            f.write(f"original_name: modified_{file_id}.xlsx\n")
            f.write(f"timestamp: {datetime.now().isoformat()}\n")
            f.write(f"type: modified\n")
        
        # Clean old revisions
        self._cleanup_old_revisions(file_id)
        
        return revision
