"""
Controlled executor for safe Excel operations using a whitelist approach.
"""
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
from typing import Dict, Any, List, Optional
import pandas as pd
import re


class ControlledExecutor:
    """Executes controlled Excel operations based on a whitelist of allowed actions."""
    
    # Whitelist of allowed operations
    ALLOWED_OPERATIONS = [
        "filter_rows",
        "sort_rows",
        "add_column",
        "remove_column",
        "rename_column",
        "fill_column",
        "replace_value",
        "sum_column",
        "average_column",
        "count_rows",
        "merge_sheets",
        "pivot_table",
        "remove_duplicates",
        "set_cell_value",
        "copy_column",
        "calculate_formula"
    ]
    
    def __init__(self):
        """Initialize the controlled executor."""
        self.operation_log = []
    
    def parse_instruction(self, instruction: str, file_metadata: Dict) -> List[Dict]:
        """Parse natural language instruction into structured operations.
        
        Args:
            instruction: Natural language instruction
            file_metadata: Metadata about the files
            
        Returns:
            List of operation dictionaries
        """
        # This is a simplified parser. In production, you'd use the LLM to generate
        # structured operations from natural language
        operations = []
        
        # Extract file references (F1, F2, etc.)
        file_refs = re.findall(r'\bF(\d+)\b', instruction)
        
        # Simple keyword matching for common operations
        instruction_lower = instruction.lower()
        
        if "filter" in instruction_lower:
            operations.append({
                "type": "filter_rows",
                "params": self._extract_filter_params(instruction)
            })
        
        if "sort" in instruction_lower:
            operations.append({
                "type": "sort_rows",
                "params": self._extract_sort_params(instruction)
            })
        
        if "add column" in instruction_lower or "new column" in instruction_lower:
            operations.append({
                "type": "add_column",
                "params": self._extract_add_column_params(instruction)
            })
        
        if "sum" in instruction_lower:
            operations.append({
                "type": "sum_column",
                "params": self._extract_aggregate_params(instruction, "sum")
            })
        
        if "average" in instruction_lower or "mean" in instruction_lower:
            operations.append({
                "type": "average_column",
                "params": self._extract_aggregate_params(instruction, "average")
            })
        
        return operations
    
    def _extract_filter_params(self, instruction: str) -> Dict:
        """Extract filter parameters from instruction."""
        # Simplified extraction
        return {
            "condition": instruction,
            "sheet": None
        }
    
    def _extract_sort_params(self, instruction: str) -> Dict:
        """Extract sort parameters from instruction."""
        return {
            "column": None,
            "ascending": "descending" not in instruction.lower(),
            "sheet": None
        }
    
    def _extract_add_column_params(self, instruction: str) -> Dict:
        """Extract add column parameters from instruction."""
        return {
            "name": "New Column",
            "formula": None,
            "sheet": None
        }
    
    def _extract_aggregate_params(self, instruction: str, operation: str) -> Dict:
        """Extract aggregate operation parameters from instruction."""
        return {
            "column": None,
            "sheet": None,
            "operation": operation
        }
    
    def execute_operations(self, operations: List[Dict], workbook: openpyxl.Workbook,
                          sheet_name: Optional[str] = None) -> Dict[str, Any]:
        """Execute a list of operations on a workbook.
        
        Args:
            operations: List of operation dictionaries
            workbook: Excel workbook object
            sheet_name: Default sheet name to operate on
            
        Returns:
            Dictionary with execution results
        """
        results = {
            "success": True,
            "operations_completed": 0,
            "errors": [],
            "log": []
        }
        
        for op in operations:
            try:
                op_type = op.get("type")
                params = op.get("params", {})
                
                if op_type not in self.ALLOWED_OPERATIONS:
                    results["errors"].append(f"Operation '{op_type}' not allowed")
                    continue
                
                # Get the target sheet
                target_sheet = params.get("sheet") or sheet_name
                if not target_sheet or target_sheet not in workbook.sheetnames:
                    target_sheet = workbook.sheetnames[0]
                
                ws = workbook[target_sheet]
                
                # Execute the operation
                method_name = f"_execute_{op_type}"
                if hasattr(self, method_name):
                    method = getattr(self, method_name)
                    method(ws, params)
                    results["operations_completed"] += 1
                    results["log"].append(f"Completed: {op_type} on {target_sheet}")
                else:
                    results["errors"].append(f"Handler not implemented for '{op_type}'")
                
            except Exception as e:
                results["errors"].append(f"Error in {op_type}: {str(e)}")
                results["success"] = False
        
        return results
    
    def _execute_set_cell_value(self, worksheet, params: Dict):
        """Set a cell value."""
        cell = params.get("cell")
        value = params.get("value")
        if cell and value is not None:
            worksheet[cell] = value
    
    def _execute_fill_column(self, worksheet, params: Dict):
        """Fill a column with a value."""
        column = params.get("column")
        value = params.get("value")
        start_row = params.get("start_row", 2)
        end_row = params.get("end_row", worksheet.max_row)
        
        if column and value is not None:
            col_letter = column if isinstance(column, str) else get_column_letter(column)
            for row in range(start_row, end_row + 1):
                worksheet[f"{col_letter}{row}"] = value
    
    def _execute_add_column(self, worksheet, params: Dict):
        """Add a new column."""
        name = params.get("name", "New Column")
        position = params.get("position", worksheet.max_column + 1)
        
        # Insert column header
        col_letter = get_column_letter(position)
        worksheet[f"{col_letter}1"] = name
    
    def _execute_replace_value(self, worksheet, params: Dict):
        """Replace values in the worksheet."""
        old_value = params.get("old_value")
        new_value = params.get("new_value")
        column = params.get("column")
        
        if old_value is not None and new_value is not None:
            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    if column:
                        col_letter = get_column_letter(cell.column)
                        if col_letter != column:
                            continue
                    if cell.value == old_value:
                        cell.value = new_value
    
    def _execute_sum_column(self, worksheet, params: Dict):
        """Calculate sum of a column."""
        column = params.get("column")
        result_cell = params.get("result_cell")
        
        if column and result_cell:
            col_letter = column if isinstance(column, str) else get_column_letter(column)
            # Find last row with data
            max_row = worksheet.max_row
            # Add formula
            worksheet[result_cell] = f"=SUM({col_letter}2:{col_letter}{max_row})"
    
    def _execute_average_column(self, worksheet, params: Dict):
        """Calculate average of a column."""
        column = params.get("column")
        result_cell = params.get("result_cell")
        
        if column and result_cell:
            col_letter = column if isinstance(column, str) else get_column_letter(column)
            max_row = worksheet.max_row
            worksheet[result_cell] = f"=AVERAGE({col_letter}2:{col_letter}{max_row})"
    
    def _execute_copy_column(self, worksheet, params: Dict):
        """Copy a column to another location."""
        source_col = params.get("source_column")
        target_col = params.get("target_column")
        
        if source_col and target_col:
            src_letter = source_col if isinstance(source_col, str) else get_column_letter(source_col)
            tgt_letter = target_col if isinstance(target_col, str) else get_column_letter(target_col)
            
            for row in range(1, worksheet.max_row + 1):
                src_cell = worksheet[f"{src_letter}{row}"]
                tgt_cell = worksheet[f"{tgt_letter}{row}"]
                tgt_cell.value = src_cell.value
    
    def execute_from_llm_response(self, llm_response: str, workbook: openpyxl.Workbook) -> Dict[str, Any]:
        """Execute operations from LLM-generated response.
        
        Args:
            llm_response: JSON string or structured response from LLM
            workbook: Excel workbook object
            
        Returns:
            Dictionary with execution results
        """
        # Parse the LLM response to extract operations
        # This would parse JSON or structured format from the LLM
        # For now, return a placeholder
        return {
            "success": False,
            "error": "LLM response parsing not yet implemented",
            "operations_completed": 0
        }
