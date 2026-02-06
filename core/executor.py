"""
Executor module for safe execution of Excel operations based on validated plans.
"""
import pandas as pd
import openpyxl
from typing import Dict, Any, List
from datetime import datetime
from .plan import (
    ExecutionPlan, Operation, OperationType,
    FilterDeleteRowsParams, DeduplicateParams, FillNullsParams,
    TypeConversionParams, ColumnSplitParams, ColumnMergeParams,
    FillStrategy, DataType
)


class ExecutionResult(BaseModel):
    """Result of executing a plan."""
    success: bool
    operations_completed: int
    operations_failed: int
    changes: List[Dict[str, Any]]  # List of changes made
    errors: List[str]
    
    def to_dict(self) -> Dict[str, Any]:
        """Convert to dictionary."""
        return {
            "success": self.success,
            "operations_completed": self.operations_completed,
            "operations_failed": self.operations_failed,
            "changes": self.changes,
            "errors": self.errors
        }


class PlanExecutor:
    """Executes validated execution plans on Excel files."""
    
    def __init__(self):
        """Initialize the executor."""
        self.execution_log = []
    
    def execute_plan(self, plan: ExecutionPlan, workbooks: Dict[str, openpyxl.Workbook]) -> Dict[str, Any]:
        """Execute a complete plan on the provided workbooks.
        
        Args:
            plan: Validated execution plan
            workbooks: Dict mapping file_alias to workbook objects
            
        Returns:
            Dictionary with execution results and change summary
        """
        results = {
            "success": True,
            "operations_completed": 0,
            "operations_failed": 0,
            "changes": [],
            "errors": []
        }
        
        for i, operation in enumerate(plan.operations):
            try:
                # Get target workbook
                wb = workbooks.get(operation.target.file_alias)
                if not wb:
                    results["errors"].append(f"Operation {i+1}: Workbook {operation.target.file_alias} not found")
                    results["operations_failed"] += 1
                    results["success"] = False
                    continue
                
                # Execute operation
                change_summary = self._execute_operation(operation, wb)
                
                if change_summary.get("success", False):
                    results["operations_completed"] += 1
                    results["changes"].append({
                        "operation": i + 1,
                        "type": operation.type.value,
                        "description": operation.description or operation.type.value,
                        **change_summary
                    })
                else:
                    results["operations_failed"] += 1
                    results["errors"].append(f"Operation {i+1}: {change_summary.get('error', 'Unknown error')}")
                    if change_summary.get("critical", False):
                        results["success"] = False
                
            except Exception as e:
                results["operations_failed"] += 1
                results["errors"].append(f"Operation {i+1}: {str(e)}")
                results["success"] = False
        
        return results
    
    def _execute_operation(self, operation: Operation, workbook: openpyxl.Workbook) -> Dict[str, Any]:
        """Execute a single operation on a workbook.
        
        Args:
            operation: The operation to execute
            workbook: The target workbook
            
        Returns:
            Dictionary with change summary
        """
        # Get worksheet
        ws = workbook[operation.target.sheet_name]
        
        # Convert worksheet to DataFrame for easier manipulation
        df = self._worksheet_to_dataframe(ws)
        initial_rows = len(df)
        initial_columns = len(df.columns)
        
        # Execute based on operation type
        if operation.type == OperationType.FILTER_DELETE_ROWS:
            result_df, summary = self._execute_filter_delete(df, operation.params)
        elif operation.type == OperationType.DEDUPLICATE:
            result_df, summary = self._execute_deduplicate(df, operation.params)
        elif operation.type == OperationType.FILL_NULLS:
            result_df, summary = self._execute_fill_nulls(df, operation.params)
        elif operation.type == OperationType.TYPE_CONVERSION:
            result_df, summary = self._execute_type_conversion(df, operation.params)
        elif operation.type == OperationType.COLUMN_SPLIT:
            result_df, summary = self._execute_column_split(df, operation.params)
        elif operation.type == OperationType.COLUMN_MERGE:
            result_df, summary = self._execute_column_merge(df, operation.params)
        else:
            return {"success": False, "error": f"Unsupported operation type: {operation.type}"}
        
        # Update summary with before/after stats
        summary["initial_rows"] = initial_rows
        summary["initial_columns"] = initial_columns
        summary["final_rows"] = len(result_df)
        summary["final_columns"] = len(result_df.columns)
        summary["rows_changed"] = initial_rows - len(result_df)
        summary["columns_changed"] = len(result_df.columns) - initial_columns
        
        # Write back to worksheet
        self._dataframe_to_worksheet(result_df, ws)
        
        return summary
    
    def _worksheet_to_dataframe(self, worksheet) -> pd.DataFrame:
        """Convert worksheet to DataFrame."""
        data = worksheet.values
        cols = next(data)
        data = list(data)
        df = pd.DataFrame(data, columns=cols)
        return df
    
    def _dataframe_to_worksheet(self, df: pd.DataFrame, worksheet):
        """Write DataFrame back to worksheet."""
        # Clear existing content
        worksheet.delete_rows(1, worksheet.max_row)
        
        # Write headers
        for col_idx, col_name in enumerate(df.columns, start=1):
            worksheet.cell(row=1, column=col_idx, value=col_name)
        
        # Write data
        for row_idx, row in enumerate(df.itertuples(index=False), start=2):
            for col_idx, value in enumerate(row, start=1):
                worksheet.cell(row=row_idx, column=col_idx, value=value)
    
    def _execute_filter_delete(self, df: pd.DataFrame, params: FilterDeleteRowsParams) -> tuple:
        """Execute filter/delete rows operation."""
        initial_count = len(df)
        column = params.column
        
        if column not in df.columns:
            return df, {"success": False, "error": f"Column {column} not found"}
        
        # Build condition mask
        if params.condition == "equals":
            mask = df[column] == params.value
        elif params.condition == "not_equals":
            mask = df[column] != params.value
        elif params.condition == "contains":
            mask = df[column].astype(str).str.contains(str(params.value), na=False)
        elif params.condition == "not_contains":
            mask = ~df[column].astype(str).str.contains(str(params.value), na=False)
        elif params.condition == "empty":
            mask = df[column].isna() | (df[column] == "")
        elif params.condition == "not_empty":
            mask = ~(df[column].isna() | (df[column] == ""))
        elif params.condition == "greater_than":
            mask = pd.to_numeric(df[column], errors='coerce') > float(params.value)
        elif params.condition == "less_than":
            mask = pd.to_numeric(df[column], errors='coerce') < float(params.value)
        else:
            return df, {"success": False, "error": f"Unknown condition: {params.condition}"}
        
        # Apply action
        if params.action == "delete":
            result_df = df[~mask].copy()
        else:  # keep
            result_df = df[mask].copy()
        
        deleted_count = initial_count - len(result_df)
        
        return result_df, {
            "success": True,
            "rows_deleted": deleted_count,
            "rows_kept": len(result_df),
            "condition": f"{params.condition} on column {column}"
        }
    
    def _execute_deduplicate(self, df: pd.DataFrame, params: DeduplicateParams) -> tuple:
        """Execute deduplication operation."""
        initial_count = len(df)
        
        # Check all columns exist
        missing_cols = [col for col in params.columns if col not in df.columns]
        if missing_cols:
            return df, {"success": False, "error": f"Columns not found: {missing_cols}"}
        
        # Remove duplicates
        result_df = df.drop_duplicates(subset=params.columns, keep=params.keep).copy()
        duplicates_removed = initial_count - len(result_df)
        
        return result_df, {
            "success": True,
            "duplicates_removed": duplicates_removed,
            "unique_rows": len(result_df),
            "checked_columns": params.columns
        }
    
    def _execute_fill_nulls(self, df: pd.DataFrame, params: FillNullsParams) -> tuple:
        """Execute fill nulls operation."""
        column = params.column
        
        if column not in df.columns:
            return df, {"success": False, "error": f"Column {column} not found"}
        
        null_count_before = df[column].isna().sum()
        result_df = df.copy()
        
        if params.strategy == FillStrategy.FIXED_VALUE:
            result_df[column] = result_df[column].fillna(params.value)
        elif params.strategy == FillStrategy.FORWARD_FILL:
            result_df[column] = result_df[column].fillna(method='ffill')
        elif params.strategy == FillStrategy.BACKWARD_FILL:
            result_df[column] = result_df[column].fillna(method='bfill')
        
        null_count_after = result_df[column].isna().sum()
        cells_filled = null_count_before - null_count_after
        
        return result_df, {
            "success": True,
            "cells_filled": int(cells_filled),
            "nulls_remaining": int(null_count_after),
            "column": column,
            "strategy": params.strategy.value
        }
    
    def _execute_type_conversion(self, df: pd.DataFrame, params: TypeConversionParams) -> tuple:
        """Execute type conversion operation."""
        column = params.column
        
        if column not in df.columns:
            return df, {"success": False, "error": f"Column {column} not found"}
        
        result_df = df.copy()
        failures = []
        success_count = 0
        
        if params.target_type == DataType.TEXT:
            result_df[column] = result_df[column].astype(str)
            success_count = len(result_df)
        
        elif params.target_type == DataType.NUMBER:
            converted = pd.to_numeric(result_df[column], errors='coerce')
            failed_mask = converted.isna() & result_df[column].notna()
            failures = result_df[failed_mask].index.tolist()
            result_df[column] = converted
            success_count = len(result_df) - len(failures)
        
        elif params.target_type in [DataType.DATE, DataType.DATETIME]:
            date_format = params.date_format if params.date_format else None
            converted = pd.to_datetime(result_df[column], format=date_format, errors='coerce')
            failed_mask = converted.isna() & result_df[column].notna()
            failures = result_df[failed_mask].index.tolist()
            result_df[column] = converted
            success_count = len(result_df) - len(failures)
        
        return result_df, {
            "success": True,
            "conversions_successful": success_count,
            "conversions_failed": len(failures),
            "failed_rows": failures[:10],  # Show first 10 failures
            "column": column,
            "target_type": params.target_type.value
        }
    
    def _execute_column_split(self, df: pd.DataFrame, params: ColumnSplitParams) -> tuple:
        """Execute column split operation."""
        column = params.source_column
        
        if column not in df.columns:
            return df, {"success": False, "error": f"Column {column} not found"}
        
        result_df = df.copy()
        
        # Split column
        max_splits = params.max_splits if params.max_splits else -1
        split_data = result_df[column].astype(str).str.split(params.delimiter, n=max_splits, expand=True)
        
        # Rename columns
        num_new_cols = min(len(params.new_column_names), split_data.shape[1])
        for i in range(num_new_cols):
            result_df[params.new_column_names[i]] = split_data[i]
        
        return result_df, {
            "success": True,
            "source_column": column,
            "new_columns": params.new_column_names[:num_new_cols],
            "columns_created": num_new_cols
        }
    
    def _execute_column_merge(self, df: pd.DataFrame, params: ColumnMergeParams) -> tuple:
        """Execute column merge operation."""
        # Check all source columns exist
        missing_cols = [col for col in params.source_columns if col not in df.columns]
        if missing_cols:
            return df, {"success": False, "error": f"Columns not found: {missing_cols}"}
        
        result_df = df.copy()
        
        # Merge columns
        result_df[params.target_column] = result_df[params.source_columns].astype(str).agg(params.delimiter.join, axis=1)
        
        # Optionally delete source columns
        columns_deleted = 0
        if params.delete_sources:
            result_df = result_df.drop(columns=params.source_columns)
            columns_deleted = len(params.source_columns)
        
        return result_df, {
            "success": True,
            "source_columns": params.source_columns,
            "target_column": params.target_column,
            "columns_deleted": columns_deleted
        }


from pydantic import BaseModel
